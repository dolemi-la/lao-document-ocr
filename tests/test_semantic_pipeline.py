from docx import Document as WordDocument

from lao_document_ocr.exporters import export_docx
from lao_document_ocr.exporters.markdown import document_to_markdown
from lao_document_ocr.models import Block, BlockType, BoundingBox, Document, Page
from lao_document_ocr.ocr.base import RecognizedLine
from lao_document_ocr.structure import build_blocks


def _line(
    text: str,
    *,
    block_id: int,
    paragraph_id: int,
    line_id: int,
    y: int,
    height: int,
) -> RecognizedLine:
    return RecognizedLine(
        text=text,
        bbox=BoundingBox(x=50, y=y, width=320, height=height),
        confidence=0.96,
        block_id=block_id,
        paragraph_id=paragraph_id,
        line_id=line_id,
    )


def test_structure_assigns_heading_level_from_relative_height() -> None:
    lines = [
        _line(
            "Main title",
            block_id=1,
            paragraph_id=1,
            line_id=1,
            y=20,
            height=40,
        ),
        _line(
            "Body one",
            block_id=2,
            paragraph_id=1,
            line_id=1,
            y=100,
            height=20,
        ),
        _line(
            "Body two",
            block_id=3,
            paragraph_id=1,
            line_id=1,
            y=150,
            height=20,
        ),
        _line(
            "Body three",
            block_id=4,
            paragraph_id=1,
            line_id=1,
            y=200,
            height=20,
        ),
    ]

    blocks = build_blocks(lines)

    assert blocks[0].type == BlockType.HEADING
    assert blocks[0].level == 1
    assert blocks[0].metadata["classifier"] == "heuristic-v2"
    assert all(block.type == BlockType.PARAGRAPH for block in blocks[1:])


def test_structure_preserves_original_list_text_but_stores_clean_items() -> None:
    lines = [
        _line(
            "1. ກາເຟ",
            block_id=1,
            paragraph_id=1,
            line_id=1,
            y=50,
            height=20,
        ),
        _line(
            "ແລະ ຊາ",
            block_id=1,
            paragraph_id=1,
            line_id=2,
            y=80,
            height=20,
        ),
        _line(
            "2. ນ້ຳ",
            block_id=1,
            paragraph_id=1,
            line_id=3,
            y=110,
            height=20,
        ),
    ]

    block = build_blocks(lines)[0]

    assert block.type == BlockType.LIST
    assert block.text == "1. ກາເຟ\nແລະ ຊາ\n2. ນ້ຳ"
    assert block.metadata["list_style"] == "ordered"
    assert block.metadata["items"] == ["ກາເຟ ແລະ ຊາ", "ນ້ຳ"]


def test_ordered_list_exports_as_word_numbered_list_without_markers(tmp_path) -> None:
    block = Block(
        type=BlockType.LIST,
        text="1. First\n2. Second",
        metadata={
            "list_style": "ordered",
            "items": ["First", "Second"],
        },
    )
    document = Document(
        pages=[
            Page(
                number=1,
                width=600,
                height=800,
                blocks=[block],
            )
        ]
    )

    path = export_docx(document, tmp_path / "numbered.docx")
    word = WordDocument(path)

    paragraphs = [paragraph for paragraph in word.paragraphs if paragraph.text]
    assert [paragraph.text for paragraph in paragraphs] == ["First", "Second"]
    assert [paragraph.style.name for paragraph in paragraphs] == [
        "List Number",
        "List Number",
    ]


def test_markdown_list_style_uses_clean_items() -> None:
    ordered = Block(
        type=BlockType.LIST,
        text="1. First\n2. Second",
        metadata={
            "list_style": "ordered",
            "items": ["First", "Second"],
        },
    )
    bullets = Block(
        type=BlockType.LIST,
        text="• Alpha\n- Beta",
        metadata={
            "list_style": "bullet",
            "items": ["Alpha", "Beta"],
        },
    )
    document = Document(
        pages=[
            Page(
                number=1,
                width=600,
                height=800,
                blocks=[ordered, bullets],
            )
        ]
    )

    markdown = document_to_markdown(document)

    assert "1. First" in markdown
    assert "2. Second" in markdown
    assert "- Alpha" in markdown
    assert "- Beta" in markdown
    assert "- 1. First" not in markdown


def test_learned_heading_hint_promotes_normal_height_text() -> None:
    lines = [
        RecognizedLine(
            text="Learned heading",
            bbox=BoundingBox(x=50, y=50, width=320, height=20),
            confidence=0.95,
            block_id=1,
            paragraph_id=1,
            line_id=1,
            semantic_type=BlockType.HEADING,
        ),
        RecognizedLine(
            text="Normal body",
            bbox=BoundingBox(x=50, y=120, width=320, height=20),
            confidence=0.95,
            block_id=2,
            paragraph_id=2,
            line_id=1,
            semantic_type=BlockType.PARAGRAPH,
        ),
    ]

    blocks = build_blocks(lines)

    assert blocks[0].type == BlockType.HEADING
    assert blocks[0].metadata["semantic_hint"] == "heading"
    assert blocks[0].metadata["semantic_hint_source"] == "learned-layout"
    assert blocks[1].type == BlockType.PARAGRAPH


def test_learned_unresolved_list_and_table_preserve_semantics_without_structure() -> None:
    lines = [
        RecognizedLine(
            text="First item",
            bbox=BoundingBox(x=50, y=50, width=220, height=20),
            confidence=0.95,
            block_id=1,
            paragraph_id=1,
            line_id=1,
            semantic_type=BlockType.LIST,
        ),
        RecognizedLine(
            text="Second item",
            bbox=BoundingBox(x=50, y=80, width=220, height=20),
            confidence=0.95,
            block_id=1,
            paragraph_id=1,
            line_id=2,
            semantic_type=BlockType.LIST,
        ),
        RecognizedLine(
            text="Unresolved table text",
            bbox=BoundingBox(x=50, y=160, width=320, height=20),
            confidence=0.95,
            block_id=2,
            paragraph_id=2,
            line_id=1,
            semantic_type=BlockType.TABLE,
        ),
    ]

    blocks = build_blocks(lines)

    assert blocks[0].type == BlockType.LIST
    assert blocks[0].metadata["list_style"] == "unresolved"
    assert blocks[0].metadata["items"] == ["First item", "Second item"]
    assert blocks[1].type == BlockType.TABLE
    assert blocks[1].cells == []
    assert blocks[1].metadata["structure_status"] == "unresolved"


def test_unresolved_learned_list_export_does_not_invent_markers(tmp_path) -> None:
    block = Block(
        type=BlockType.LIST,
        text="First item\nSecond item",
        metadata={
            "list_style": "unresolved",
            "items": ["First item", "Second item"],
            "semantic_hint": "list",
        },
    )
    document = Document(
        pages=[Page(number=1, width=600, height=800, blocks=[block])]
    )

    path = export_docx(document, tmp_path / "unresolved-list.docx")
    word = WordDocument(path)
    paragraphs = [paragraph for paragraph in word.paragraphs if paragraph.text]
    assert [paragraph.text for paragraph in paragraphs] == [
        "First item",
        "Second item",
    ]
    assert all(paragraph.style.name == "Normal" for paragraph in paragraphs)

    markdown = document_to_markdown(document)
    assert "First item" in markdown
    assert "Second item" in markdown
    assert "- First item" not in markdown
    assert "1. First item" not in markdown
