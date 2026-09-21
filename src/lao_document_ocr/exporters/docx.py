from __future__ import annotations

import base64
import io
from pathlib import Path

from docx import Document as WordDocument
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from lao_document_ocr.header_footer import repeated_role_texts
from lao_document_ocr.models import Block, BlockType, Document

DEFAULT_FONT = "Noto Sans Lao"
DEFAULT_LANGUAGE = "lo-LA"


def _set_run_language(run, language: str = DEFAULT_LANGUAGE) -> None:
    run_properties = run._element.get_or_add_rPr()
    lang = run_properties.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        run_properties.append(lang)
    lang.set(qn("w:val"), language)
    lang.set(qn("w:eastAsia"), language)
    lang.set(qn("w:bidi"), language)


def _set_run_font(run, font_name: str) -> None:
    run.font.name = font_name
    run_properties = run._element.get_or_add_rPr()
    fonts = run_properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, fonts)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), font_name)


def _add_text(paragraph, text: str, font_name: str, bold: bool = False) -> None:
    parts = text.splitlines() or [text]
    for index, part in enumerate(parts):
        run = paragraph.add_run(part)
        run.bold = bold
        _set_run_font(run, font_name)
        _set_run_language(run)
        if index < len(parts) - 1:
            run.add_break(WD_BREAK.LINE)


def _add_block(word: WordDocument, block: Block, font_name: str) -> None:
    if block.type == BlockType.HEADING:
        paragraph = word.add_heading(level=max(1, min(9, block.level or 1)))
        _add_text(paragraph, block.text, font_name, bold=True)
        return

    if block.type == BlockType.LIST:
        style = (
            "List Number"
            if block.metadata.get("list_style") == "ordered"
            else "List Bullet"
        )
        items = block.metadata.get("items")
        if not isinstance(items, list) or not items:
            items = block.text.splitlines()
        for item in items:
            paragraph = word.add_paragraph(style=style)
            _add_text(paragraph, str(item), font_name)
        return

    if block.type == BlockType.IMAGE:
        encoded = block.metadata.get("image_base64")
        if not isinstance(encoded, str) or not encoded:
            return
        try:
            data = base64.b64decode(encoded, validate=True)
            ratio = float(block.metadata.get("width_ratio", 0.5))
        except (ValueError, TypeError):
            return
        width_inches = max(0.5, min(6.5, 6.5 * max(0.05, min(1.0, ratio))))
        try:
            word.add_picture(io.BytesIO(data), width=Inches(width_inches))
        except Exception:
            return
        return

    if block.type == BlockType.TABLE and block.cells:
        row_count = int(block.metadata.get("rows") or (max(cell.row for cell in block.cells) + 1))
        column_count = int(
            block.metadata.get("columns")
            or (max(cell.column for cell in block.cells) + 1)
        )
        table = word.add_table(rows=row_count, cols=column_count)
        table.style = "Table Grid"
        for cell in sorted(block.cells, key=lambda item: (item.row, item.column)):
            target = table.cell(cell.row, cell.column)
            if cell.row_span > 1 or cell.column_span > 1:
                end_row = min(row_count - 1, cell.row + cell.row_span - 1)
                end_column = min(column_count - 1, cell.column + cell.column_span - 1)
                target = target.merge(table.cell(end_row, end_column))
            target.text = ""
            _add_text(target.paragraphs[0], cell.text, font_name)
        return

    paragraph = word.add_paragraph()
    _add_text(paragraph, block.text, font_name)


def _apply_headers_footers(
    word: WordDocument,
    document: Document,
    font_name: str,
) -> None:
    if not word.sections:
        return

    section = word.sections[0]
    header_texts = repeated_role_texts(document.pages, "header")
    footer_texts = repeated_role_texts(document.pages, "footer")

    if header_texts:
        paragraph = section.header.paragraphs[0]
        paragraph.text = ""
        _add_text(paragraph, "\n".join(header_texts), font_name)

    if footer_texts:
        paragraph = section.footer.paragraphs[0]
        paragraph.text = ""
        _add_text(paragraph, "\n".join(footer_texts), font_name)


def export_docx(
    document: Document,
    path: str | Path,
    *,
    font_name: str = DEFAULT_FONT,
) -> Path:
    destination = Path(path)
    word = WordDocument()

    normal_style = word.styles["Normal"]
    normal_style.font.name = font_name
    normal_style.font.size = Pt(11)

    _apply_headers_footers(word, document, font_name)

    for page_index, page in enumerate(document.pages):
        for block in page.blocks:
            if block.metadata.get("role") in {"header", "footer"}:
                continue
            _add_block(word, block, font_name)
        if page_index < len(document.pages) - 1:
            word.add_page_break()

    core = word.core_properties
    core.title = document.source_name or "OCR document"
    core.subject = "Generated by Lao Document OCR"

    word.save(destination)
    return destination
